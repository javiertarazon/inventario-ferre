"""Tests for reports blueprint including SENIAT operational report."""

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.extensions import db
from app.models import DailySalesClosure, ExchangeRate, ItemGroup, Movimiento, Product, PurchaseInvoice, Proveedor



def login_test_user(client):
    return client.post(
        '/login',
        data={'username': 'testuser', 'password': 'testpass123'},
        follow_redirects=False,
    )



def test_seniat_operativo_report_renders_purchases_and_closures(app, client, test_user):
    """The SENIAT report should show purchase invoices and daily closures for the selected period."""
    user_id, _ = test_user

    with app.app_context():
        supplier = Proveedor(
            nombre='Proveedor Fiscal',
            rif='J-99999999-1',
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add(supplier)
        db.session.flush()

        db.session.add_all([
            ExchangeRate(date=date(2025, 3, 5), rate=Decimal('36.5000'), created_by=user_id),
            ExchangeRate(date=date(2025, 3, 10), rate=Decimal('38.0000'), created_by=user_id),
        ])
        db.session.add(PurchaseInvoice(
            supplier_id=supplier.id,
            invoice_number='FAC-SEN-001',
            invoice_date=date(2025, 3, 5),
            currency_code='USD',
            exchange_rate_value=Decimal('36.5000'),
            subtotal_usd=Decimal('100.0000'),
            taxable_base_usd=Decimal('100.0000'),
            tax_amount_usd=Decimal('16.0000'),
            total_usd=Decimal('116.0000'),
            created_by=user_id,
            updated_by=user_id,
        ))
        db.session.add(DailySalesClosure(
            closure_date=date(2025, 3, 10),
            total_sales_usd=Decimal('10.0000'),
            invoiced_share=Decimal('0.6000'),
            non_invoiced_share=Decimal('0.4000'),
            invoiced_sales_usd=Decimal('6.0000'),
            non_invoiced_sales_usd=Decimal('4.0000'),
            reconstructed_sales_usd=Decimal('10.0000'),
            generated_exit_units=5,
            source_file_name='cierres_marzo.xlsx',
            created_by=user_id,
            updated_by=user_id,
        ))
        db.session.commit()

    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/seniat?start_date=2025-03-01&end_date=2025-03-31')
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'Reporte SENIAT Operativo' in html
    assert 'FAC-SEN-001' in html
    assert 'Proveedor Fiscal' in html
    assert '116.0000' in html
    assert '4234.00' in html
    assert '380.00' in html



def test_seniat_operativo_export_returns_excel(app, client, test_user):
    """The SENIAT report export should return an Excel workbook."""
    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/seniat/export?start_date=2025-03-01&end_date=2025-03-31')

    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'reporte_seniat_operativo_2025-03-01_2025-03-31.xlsx' in response.headers['Content-Disposition']



def test_libro_inventario_export_omits_precio_usd_column(app, client, test_user):
    """The Excel export of the inventory book must not include the Precio USD column."""
    user_id, _ = test_user

    with app.app_context():
        category = ItemGroup(
            name='Categoria SENIAT',
            description='Categoria para prueba de libro inventario',
            color='#0d6efd',
            icon='bi-box',
            created_by=user_id,
            updated_by=user_id,
        )
        supplier = Proveedor(
            nombre='Proveedor Libro',
            rif='J-11111111-1',
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add_all([
            category,
            supplier,
            ExchangeRate(date=date(2025, 3, 7), rate=Decimal('40.0000'), created_by=user_id),
        ])
        db.session.flush()

        db.session.add(Product(
            codigo='L-IV-01',
            descripcion='Producto libro inventario',
            stock=5,
            precio_dolares=Decimal('10.00'),
            factor_ajuste=Decimal('1.00'),
            proveedor_id=supplier.id,
            item_group_id=category.id,
            created_by=user_id,
            updated_by=user_id,
        ))
        db.session.commit()

    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/libro-inventario/export?fecha=2025-03-07')

    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    workbook = load_workbook(io.BytesIO(response.data), data_only=True)
    worksheet = workbook['Libro Inventario']
    headers = [cell for cell in next(worksheet.iter_rows(min_row=6, max_row=6, values_only=True))]

    assert 'Precio USD' not in headers
    assert headers == [
        'Codigo',
        'Descripcion',
        'Proveedor',
        'Stock',
        'Precio Bs',
        'Valor Total USD',
        'Valor Total Bs',
    ]



def test_movimientos_export_uses_bs_only_and_daily_rate(app, client, test_user):
    """Movements Excel export must follow the SENIAT visual template and value rows with the exact rate of each day."""
    user_id, _ = test_user

    with app.app_context():
        category = ItemGroup(
            name='Categoria Movimientos',
            description='Categoria de prueba movimientos',
            color='#198754',
            icon='bi-box',
            created_by=user_id,
            updated_by=user_id,
        )
        supplier = Proveedor(
            nombre='Proveedor Movimientos',
            rif='J-22222222-2',
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add_all([
            category,
            supplier,
            ExchangeRate(date=date(2025, 3, 5), rate=Decimal('10.0000'), created_by=user_id),
            ExchangeRate(date=date(2025, 3, 6), rate=Decimal('20.0000'), created_by=user_id),
        ])
        db.session.flush()

        product = Product(
            codigo='M-EX-01',
            descripcion='Producto movimientos export',
            stock=10,
            precio_dolares=Decimal('2.00'),
            factor_ajuste=Decimal('1.00'),
            proveedor_id=supplier.id,
            item_group_id=category.id,
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add(product)
        db.session.flush()

        db.session.add_all([
            Movimiento(
                producto_id=product.id,
                tipo='SALIDA',
                cantidad=1,
                fecha=date(2025, 3, 5),
                descripcion='Movimiento dia 1',
                created_at=datetime(2025, 3, 5, 9, 0, 0),
                created_by=user_id,
                updated_by=user_id,
            ),
            Movimiento(
                producto_id=product.id,
                tipo='SALIDA',
                cantidad=2,
                fecha=date(2025, 3, 6),
                descripcion='Movimiento dia 2',
                created_at=datetime(2025, 3, 6, 10, 0, 0),
                created_by=user_id,
                updated_by=user_id,
            ),
        ])
        db.session.commit()

    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/movimientos/export?start_date=2025-03-05&end_date=2025-03-06')

    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.data), data_only=True)
    worksheet = workbook['Movimientos']
    group_headers = [cell for cell in next(worksheet.iter_rows(min_row=9, max_row=9, values_only=True))]
    headers = [cell for cell in next(worksheet.iter_rows(min_row=10, max_row=10, values_only=True))]
    row_1 = [cell for cell in next(worksheet.iter_rows(min_row=11, max_row=11, values_only=True))]
    row_2 = [cell for cell in next(worksheet.iter_rows(min_row=12, max_row=12, values_only=True))]

    assert worksheet['A1'].value == 'EMPRESA '
    assert worksheet['A5'].value == 'Movimiento de Unidades según el artículo 177 ley de impuesto  sobre la renta '
    assert group_headers == [
        '  ', None, None,
        ' Existencia Inicial', None, None,
        ' Entradas', None, None,
        ' Salidas', None, None,
        ' Autoconsumos', None, None,
        ' Inv. Actual', None,
    ]
    assert headers == [
        ' Código', ' Descripción', ' Costo Unitario', ' Cantidad', ' Monto',
        ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
        ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
    ]
    assert row_1[0] == 'M-EX-01'
    assert row_1[8] == 40
    assert row_1[9] == 2
    assert row_1[10] == 80
    assert row_2[0] == 'M-EX-01'
    assert row_2[8] == 20
    assert row_2[9] == 1
    assert row_2[10] == 20



def test_libro_diario_export_uses_bs_only_and_daily_rate(app, client, test_user):
    """Daily journal export must use the SENIAT visual template and the movement date exchange rate in Bs."""
    user_id, _ = test_user

    with app.app_context():
        category = ItemGroup(
            name='Categoria Libro Diario',
            description='Categoria de prueba libro diario',
            color='#fd7e14',
            icon='bi-box',
            created_by=user_id,
            updated_by=user_id,
        )
        supplier = Proveedor(
            nombre='Proveedor Libro Diario',
            rif='J-33333333-3',
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add_all([
            category,
            supplier,
            ExchangeRate(date=date(2025, 3, 5), rate=Decimal('10.0000'), created_by=user_id),
            ExchangeRate(date=date(2025, 3, 6), rate=Decimal('20.0000'), created_by=user_id),
        ])
        db.session.flush()

        product = Product(
            codigo='L-DI-01',
            descripcion='Producto libro diario',
            stock=10,
            precio_dolares=Decimal('2.00'),
            factor_ajuste=Decimal('1.00'),
            proveedor_id=supplier.id,
            item_group_id=category.id,
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add(product)
        db.session.flush()

        db.session.add_all([
            Movimiento(
                producto_id=product.id,
                tipo='SALIDA',
                cantidad=1,
                fecha=date(2025, 3, 5),
                descripcion='Salida dia 1',
                created_at=datetime(2025, 3, 5, 8, 30, 0),
                created_by=user_id,
                updated_by=user_id,
            ),
            Movimiento(
                producto_id=product.id,
                tipo='ENTRADA',
                cantidad=2,
                fecha=date(2025, 3, 6),
                descripcion='Entrada dia 2',
                created_at=datetime(2025, 3, 6, 11, 15, 0),
                created_by=user_id,
                updated_by=user_id,
            ),
        ])
        db.session.commit()

    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/libro-diario/export?start_date=2025-03-05&end_date=2025-03-06')

    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.data), data_only=True)
    asientos = workbook['Asientos']
    resumen = workbook['Resumen']
    asientos_headers = [cell for cell in next(asientos.iter_rows(min_row=10, max_row=10, values_only=True))]
    resumen_headers = [cell for cell in next(resumen.iter_rows(min_row=10, max_row=10, values_only=True))]
    asientos_row_1 = [cell for cell in next(asientos.iter_rows(min_row=11, max_row=11, values_only=True))]
    asientos_row_2 = [cell for cell in next(asientos.iter_rows(min_row=12, max_row=12, values_only=True))]
    resumen_row_1 = [cell for cell in next(resumen.iter_rows(min_row=11, max_row=11, values_only=True))]

    assert asientos_headers == resumen_headers
    assert asientos_row_1[0] == 'L-DI-01'
    assert asientos_row_1[8] == 20
    assert asientos_row_1[9] == 1
    assert asientos_row_1[10] == 20
    assert asientos_row_2[0] == 'L-DI-01'
    assert asientos_row_2[5] == 40
    assert asientos_row_2[6] == 2
    assert asientos_row_2[7] == 80
    assert resumen_row_1[2] == 20
    assert resumen_row_1[3] == 9
    assert resumen_row_1[4] == 180
    assert resumen_row_1[5] == 40
    assert resumen_row_1[6] == 2
    assert resumen_row_1[7] == 80
    assert resumen_row_1[8] == 20
    assert resumen_row_1[9] == 1
    assert resumen_row_1[10] == 20
    assert resumen_row_1[14] == 40
    assert resumen_row_1[15] == 10
    assert resumen_row_1[16] == 400



def test_resumen_mensual_export_uses_bs_only_and_period_rates(app, client, test_user):
    """Monthly summary export must follow the SENIAT visual template and use opening/closing rates of the period."""
    user_id, _ = test_user

    with app.app_context():
        category = ItemGroup(
            name='Categoria Mensual',
            description='Categoria de prueba resumen mensual',
            color='#6f42c1',
            icon='bi-box',
            created_by=user_id,
            updated_by=user_id,
        )
        supplier = Proveedor(
            nombre='Proveedor Mensual',
            rif='J-44444444-4',
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add_all([
            category,
            supplier,
            ExchangeRate(date=date(2025, 3, 1), rate=Decimal('5.0000'), created_by=user_id),
            ExchangeRate(date=date(2025, 3, 31), rate=Decimal('15.0000'), created_by=user_id),
        ])
        db.session.flush()

        product = Product(
            codigo='R-ME-01',
            descripcion='Producto resumen mensual',
            stock=10,
            precio_dolares=Decimal('2.00'),
            factor_ajuste=Decimal('1.00'),
            proveedor_id=supplier.id,
            item_group_id=category.id,
            created_by=user_id,
            updated_by=user_id,
        )
        db.session.add(product)
        db.session.flush()

        db.session.add(Movimiento(
            producto_id=product.id,
            tipo='SALIDA',
            cantidad=2,
            fecha=date(2025, 3, 15),
            descripcion='Salida mensual',
            created_at=datetime(2025, 3, 15, 12, 0, 0),
            created_by=user_id,
            updated_by=user_id,
        ))
        db.session.commit()

    login_response = login_test_user(client)
    assert login_response.status_code == 302

    response = client.get('/reports/resumen-mensual/export?year=2025&month=3')

    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.data), data_only=True)
    worksheet = workbook['Resumen Mensual']
    headers = [cell for cell in next(worksheet.iter_rows(min_row=10, max_row=10, values_only=True))]
    row_1 = None
    for row in worksheet.iter_rows(min_row=11, values_only=True):
        if row[0] == 'R-ME-01':
            row_1 = list(row)
            break

    assert headers == [
        ' Código', ' Descripción', ' Costo Unitario', ' Cantidad', ' Monto',
        ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
        ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
    ]
    assert row_1 is not None
    assert row_1[2] == 10
    assert row_1[3] == 12
    assert row_1[4] == 120
    assert row_1[8] == 10
    assert row_1[9] == 2
    assert row_1[10] == 20
    assert row_1[14] == 30
    assert row_1[15] == 10
    assert row_1[16] == 300
