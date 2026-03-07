"""
Reports blueprint - Routes for inventory and sales reports.
"""
import io
from datetime import datetime, date, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from flask import (
    Blueprint, render_template, redirect, url_for,
    request, flash, send_file, current_app
)
from flask_login import login_required

from app.services.company_settings_service import CompanySettingsService
from app.services.reports_service import ReportsService

reports_bp = Blueprint('reports', __name__)


SENIAT_TEMPLATE_TITLE = 'Movimiento de Unidades según el artículo 177 ley de impuesto  sobre la renta '
SENIAT_COLUMN_WIDTHS = {
    'A': 17.140625,
    'B': 42.85546875,
    'C': 13.5703125,
    'D': 15.7109375,
    'E': 11.28515625,
    'F': 13.5703125,
    'G': 9.28515625,
    'H': 11.5703125,
    'I': 13.5703125,
    'J': 9.28515625,
    'K': 9.0,
    'L': 13.5703125,
    'M': 14.0,
    'N': 7.140625,
    'O': 13.5703125,
    'P': 10.5703125,
    'Q': 10.85546875,
}
SENIAT_GROUP_HEADERS = {
    'A9': '  ',
    'D9': ' Existencia Inicial',
    'G9': ' Entradas',
    'J9': ' Salidas',
    'M9': ' Autoconsumos',
    'P9': ' Inv. Actual',
}
SENIAT_SUBHEADERS = [
    ' Código', ' Descripción', ' Costo Unitario', ' Cantidad', ' Monto',
    ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
    ' Costo Unitario', ' Cantidad', ' Monto', ' Costo Unitario', ' Cantidad', ' Monto',
]
SENIAT_NUMERIC_COLUMNS = {3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17}
SENIAT_MONEY_COLUMNS = {3, 5, 6, 8, 9, 11, 12, 14, 15, 17}
SENIAT_QTY_COLUMNS = {4, 7, 10, 13, 16}


def _seniat_border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _safe_unit_cost(total_amount, quantity, fallback=0):
    if quantity:
        return round(total_amount / quantity, 2)
    return round(fallback or 0, 2)


def _build_seniat_row(
    codigo,
    descripcion,
    costo_inicial=0,
    cantidad_inicial=0,
    monto_inicial=0,
    costo_entrada=0,
    cantidad_entrada=0,
    monto_entrada=0,
    costo_salida=0,
    cantidad_salida=0,
    monto_salida=0,
    costo_autoconsumo=0,
    cantidad_autoconsumo=0,
    monto_autoconsumo=0,
    costo_actual=0,
    cantidad_actual=0,
    monto_actual=0,
):
    return [
        codigo,
        descripcion,
        round(costo_inicial or 0, 2),
        cantidad_inicial or 0,
        round(monto_inicial or 0, 2),
        round(costo_entrada or 0, 2),
        cantidad_entrada or 0,
        round(monto_entrada or 0, 2),
        round(costo_salida or 0, 2),
        cantidad_salida or 0,
        round(monto_salida or 0, 2),
        round(costo_autoconsumo or 0, 2),
        cantidad_autoconsumo or 0,
        round(monto_autoconsumo or 0, 2),
        round(costo_actual or 0, 2),
        cantidad_actual or 0,
        round(monto_actual or 0, 2),
    ]


def _write_seniat_visual_sheet(workbook, sheet_name, company, start_date, end_date, rows, title=None):
    """Write one worksheet aligned to the reference SENIAT visual template."""
    ws = workbook.create_sheet(title=sheet_name)
    border = _seniat_border()

    ws['A1'] = 'EMPRESA '
    ws['B1'] = company['company_name']
    ws['A2'] = 'R.I.F.'
    ws['B2'] = f" {company['rif']}"
    ws['A3'] = 'DIRECCION '
    ws['B3'] = company.get('fiscal_address', '')
    ws['A4'] = 'TELEFONO '
    ws['B4'] = company.get('phone', '')
    ws['A5'] = title or SENIAT_TEMPLATE_TITLE
    ws['A6'] = f"Fecha Desde : {start_date.strftime('%d/%m/%Y')}"
    ws['A7'] = f"Fecha Hasta : {end_date.strftime('%d/%m/%Y')}"

    for coord, value in SENIAT_GROUP_HEADERS.items():
        ws[coord] = value
    for column_index, value in enumerate(SENIAT_SUBHEADERS, start=1):
        ws.cell(row=10, column=column_index, value=value)

    for column_letter, width in SENIAT_COLUMN_WIDTHS.items():
        ws.column_dimensions[column_letter].width = width

    for row_index in range(1, 8):
        for column_index in range(1, 18):
            cell = ws.cell(row=row_index, column=column_index)
            cell.font = Font(size=10)

    for row_index in (9, 10):
        for column_index in range(1, 18):
            cell = ws.cell(row=row_index, column=column_index)
            cell.font = Font(size=10)
            cell.border = border
            if row_index == 9:
                if column_index == 1:
                    cell.alignment = Alignment(horizontal='right')
                elif column_index in {4, 7, 10, 13, 16}:
                    cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(horizontal='center')

    for row_offset, row_values in enumerate(rows, start=11):
        for column_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_offset, column=column_index, value=value)
            cell.font = Font(size=10)
            cell.border = border
            if column_index == 2:
                cell.alignment = Alignment(horizontal='left', wrap_text=True)
            elif column_index in SENIAT_NUMERIC_COLUMNS:
                cell.alignment = Alignment(horizontal='right')
                if column_index in SENIAT_MONEY_COLUMNS:
                    cell.number_format = '0.00'
                elif column_index in SENIAT_QTY_COLUMNS:
                    cell.number_format = '0.####'
            else:
                cell.alignment = Alignment(horizontal='left')

    ws.freeze_panes = 'A11'
    return ws


def _build_seniat_workbook(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        _write_seniat_visual_sheet(
            workbook,
            sheet_name=sheet['sheet_name'],
            company=sheet['company'],
            start_date=sheet['start_date'],
            end_date=sheet['end_date'],
            rows=sheet['rows'],
            title=sheet.get('title'),
        )
    return workbook


# ---------------------------------------------------------------------------
# Index
# ---------------------------------------------------------------------------

@reports_bp.route('/')
@login_required
def index():
    """Reports landing page."""
    return render_template('reports/index.html')


@reports_bp.route('/seniat')
@login_required
def seniat_operativo():
    """Operational fiscal report based on purchase invoices and daily closures."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)

        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())

        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        data = service.get_reporte_seniat_operativo(start_date, end_date)
        return render_template(
            'reports/seniat_operativo.html',
            data=data,
            start_date=start_date,
            end_date=end_date,
        )
    except Exception as e:
        flash(f'Error al generar reporte SENIAT: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/seniat/export')
@login_required
def seniat_operativo_export():
    """Export operational fiscal report to Excel."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)

        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())

        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        data = service.get_reporte_seniat_operativo(start_date, end_date)
        compras_df = pd.DataFrame([{
            'Fecha': row['fecha'].strftime('%Y-%m-%d') if row['fecha'] else '',
            'Factura': row['documento'],
            'Proveedor': row['proveedor'],
            'Moneda': row['moneda'],
            'Tasa BCV': row['tasa'],
            'Base USD': row['base_usd'],
            'Impuesto USD': row['impuesto_usd'],
            'Total USD': row['total_usd'],
            'Total Bs': row['total_bs'],
            'Notas': row['notas'],
        } for row in data['compras']])
        cierres_df = pd.DataFrame([{
            'Fecha': row['fecha'].strftime('%Y-%m-%d') if row['fecha'] else '',
            'Tasa BCV': row['tasa'],
            'Total USD': row['total_usd'],
            'Total Bs': row['total_bs'],
            'Con factura USD': row['con_factura_usd'],
            'Sin factura USD': row['sin_factura_usd'],
            'Reconstruido USD': row['reconstruido_usd'],
            'Unidades salida': row['unidades_salida'],
            'Origen': row['origen'],
            'Notas': row['notas'],
        } for row in data['cierres']])
        resumen_df = pd.DataFrame([{
            'Compras registradas': data['resumen']['compras_count'],
            'Compras total USD': data['resumen']['compras_total_usd'],
            'Compras total Bs': data['resumen']['compras_total_bs'],
            'Cierres registrados': data['resumen']['cierres_count'],
            'Cierres total USD': data['resumen']['cierres_total_usd'],
            'Cierres total Bs': data['resumen']['cierres_total_bs'],
            'Unidades salida': data['resumen']['cierres_unidades_salida'],
        }])

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            compras_df.to_excel(writer, index=False, sheet_name='Compras')
            cierres_df.to_excel(writer, index=False, sheet_name='Cierres')
            resumen_df.to_excel(writer, index=False, sheet_name='Resumen')
        output.seek(0)

        filename = f'reporte_seniat_operativo_{start_date}_{end_date}.xlsx'
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar reporte SENIAT: {str(e)}', 'error')
        return redirect(url_for('reports.seniat_operativo'))


# ---------------------------------------------------------------------------
# Inventario completo
# ---------------------------------------------------------------------------

@reports_bp.route('/inventario')
@login_required
def inventario():
    """Full inventory report."""
    try:
        service = ReportsService()
        items = service.get_inventario_completo()
        total_usd = sum(i['valor_total_usd'] for i in items)
        total_bs = sum(i['valor_total_bs'] for i in items)
        return render_template(
            'reports/inventario.html',
            items=items,
            total_usd=round(total_usd, 2),
            total_bs=round(total_bs, 2),
            today=date.today(),
        )
    except Exception as e:
        flash(f'Error al generar reporte: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/inventario/export')
@login_required
def inventario_export():
    """Export full inventory to Excel."""
    try:
        service = ReportsService()
        items = service.get_inventario_completo()
        df = pd.DataFrame([{
            'Codigo': i['codigo'],
            'Descripcion': i['descripcion'],
            'Categoria': i['categoria'],
            'Proveedor': i['proveedor'],
            'Stock': i['stock'],
            'Precio USD': i['precio_dolares'],
            'Factor Ajuste': i['factor_ajuste'],
            'Precio Ajustado USD': i['precio_ajustado_usd'],
            'Precio Bs': i['precio_bs'],
            'Valor Total USD': i['valor_total_usd'],
            'Valor Total Bs': i['valor_total_bs'],
        } for i in items])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Inventario')
        output.seek(0)
        filename = f"inventario_completo_{date.today().strftime('%Y-%m-%d')}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.inventario'))


# ---------------------------------------------------------------------------
# Stock bajo
# ---------------------------------------------------------------------------

@reports_bp.route('/stock-bajo')
@login_required
def stock_bajo():
    """Low stock report."""
    try:
        service = ReportsService()
        items = service.get_stock_bajo()
        return render_template(
            'reports/stock_bajo.html',
            items=items,
            today=date.today(),
        )
    except Exception as e:
        flash(f'Error al generar reporte: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/stock-bajo/export')
@login_required
def stock_bajo_export():
    """Export low stock report to Excel."""
    try:
        service = ReportsService()
        items = service.get_stock_bajo()
        df = pd.DataFrame([{
            'Codigo': i['codigo'],
            'Descripcion': i['descripcion'],
            'Categoria': i['categoria'],
            'Proveedor': i['proveedor'],
            'Stock Actual': i['stock'],
            'Punto de Reorden': i['reorder_point'],
            'Cantidad a Pedir': i['reorder_quantity'],
            'Deficit': i['deficit'],
            'Estado': 'AGOTADO' if i['agotado'] else 'STOCK BAJO',
        } for i in items])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Stock Bajo')
        output.seek(0)
        filename = f"stock_bajo_{date.today().strftime('%Y-%m-%d')}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.stock_bajo'))


# ---------------------------------------------------------------------------
# Movimientos por periodo
# ---------------------------------------------------------------------------

@reports_bp.route('/movimientos')
@login_required
def movimientos():
    """Movement report with date range filter."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)

        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())
        tipo = request.args.get('tipo', '')

        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        items = service.get_movimientos_periodo(start_date, end_date, tipo or None)
        resumen = service.get_resumen_movimientos(start_date, end_date)

        return render_template(
            'reports/movimientos.html',
            items=items,
            resumen=resumen,
            start_date=start_date,
            end_date=end_date,
            tipo=tipo,
        )
    except Exception as e:
        flash(f'Error al generar reporte: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/movimientos/export')
@login_required
def movimientos_export():
    """Export movements report to Excel."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)
        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())
        tipo = request.args.get('tipo', '')
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        items = service.get_movimientos_periodo(start_date, end_date, tipo or None)
        rows = []
        for item in items:
            es_autoconsumo = item['tipo'].upper() == 'SALIDA' and 'autoconsumo' in (item['descripcion'] or '').lower()
            descripcion = (
                f"{item['producto_descripcion']} | "
                f"{item['fecha'].strftime('%d/%m/%Y') if item['fecha'] else ''} | "
                f"{item['tipo']} | {item['descripcion'] or 'Sin detalle'} | {item['usuario']}"
            )
            rows.append(_build_seniat_row(
                codigo=item['producto_codigo'],
                descripcion=descripcion,
                costo_inicial=item['precio_bs'],
                costo_entrada=item['precio_bs'] if item['tipo'].upper() == 'ENTRADA' else 0,
                cantidad_entrada=item['cantidad'] if item['tipo'].upper() == 'ENTRADA' else 0,
                monto_entrada=item['valor_bs'] if item['tipo'].upper() == 'ENTRADA' else 0,
                costo_salida=item['precio_bs'] if item['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                cantidad_salida=item['cantidad'] if item['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                monto_salida=item['valor_bs'] if item['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                costo_autoconsumo=item['precio_bs'] if es_autoconsumo else 0,
                cantidad_autoconsumo=item['cantidad'] if es_autoconsumo else 0,
                monto_autoconsumo=item['valor_bs'] if es_autoconsumo else 0,
                costo_actual=item['precio_bs'],
                cantidad_actual=item.get('stock_actual', 0),
                monto_actual=item.get('valor_actual_bs', 0),
            ))
        output = io.BytesIO()
        company = CompanySettingsService().get_company_context()
        workbook = _build_seniat_workbook([{
            'sheet_name': 'Movimientos',
            'company': company,
            'start_date': start_date,
            'end_date': end_date,
            'rows': rows,
            'title': SENIAT_TEMPLATE_TITLE,
        }])
        workbook.save(output)
        output.seek(0)
        filename = f"movimientos_{start_date}_{end_date}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.movimientos'))


# ---------------------------------------------------------------------------
# Ventas por periodo
# ---------------------------------------------------------------------------

@reports_bp.route('/ventas')
@login_required
def ventas():
    """Sales report with date range filter."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)

        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())
        status_filter = request.args.get('status', '')

        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        orders = service.get_ventas_periodo(start_date, end_date, status_filter or None)
        resumen = service.get_resumen_ventas(start_date, end_date)

        return render_template(
            'reports/ventas.html',
            orders=orders,
            resumen=resumen,
            start_date=start_date,
            end_date=end_date,
            status_filter=status_filter,
        )
    except Exception as e:
        flash(f'Error al generar reporte: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/ventas/export')
@login_required
def ventas_export():
    """Export sales report to Excel."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)
        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())
        status_filter = request.args.get('status', '')
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        orders = service.get_ventas_periodo(start_date, end_date, status_filter or None)
        df = pd.DataFrame([{
            'N. Orden': o['order_number'],
            'Fecha': o['order_date'].strftime('%Y-%m-%d') if o['order_date'] else '',
            'Cliente': o['customer'],
            'Estado': o['status'],
            'Estado Pago': o['payment_status'],
            'Total USD': o['total_amount'],
            'Cobrado USD': o['paid_amount'],
            'Pendiente USD': o['pending_amount'],
        } for o in orders])
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Ventas')
        output.seek(0)
        filename = f"ventas_{start_date}_{end_date}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.ventas'))


# ---------------------------------------------------------------------------
# Productos mas vendidos
# ---------------------------------------------------------------------------

@reports_bp.route('/top-productos')
@login_required
def top_productos():
    """Top selling products report."""
    try:
        service = ReportsService()
        default_start, default_end = service.get_default_date_range(30)

        start_str = request.args.get('start_date', default_start.isoformat())
        end_str = request.args.get('end_date', default_end.isoformat())
        limit = request.args.get('limit', 20, type=int)

        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date, end_date = default_start, default_end

        items = service.get_top_productos(limit=limit, start_date=start_date, end_date=end_date)
        return render_template(
            'reports/top_productos.html',
            items=items,
            start_date=start_date,
            end_date=end_date,
            limit=limit,
        )
    except Exception as e:
        flash(f'Error al generar reporte: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


# ---------------------------------------------------------------------------
# Libro de Inventario (Art. 177 ISLR)
# ---------------------------------------------------------------------------

@reports_bp.route('/libro-inventario')
@login_required
def libro_inventario():
    """Libro de Inventario diario (Art. 177 ISLR Venezuela)."""
    try:
        fecha_str = request.args.get('fecha', date.today().isoformat())
        try:
            fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
        except ValueError:
            fecha = date.today()

        service = ReportsService()
        data = service.get_libro_inventario(fecha)
        return render_template('reports/libro_inventario.html', data=data, fecha=fecha)
    except Exception as e:
        flash(f'Error al generar Libro de Inventario: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/libro-inventario/export')
@login_required
def libro_inventario_export():
    """Export Libro de Inventario to Excel."""
    try:
        fecha_str = request.args.get('fecha', date.today().isoformat())
        try:
            fecha = datetime.strptime(fecha_str, '%Y-%m-%d').date()
        except ValueError:
            fecha = date.today()

        service = ReportsService()
        data = service.get_libro_inventario(fecha)

        df = pd.DataFrame([{
            'Codigo': i['codigo'],
            'Descripcion': i['descripcion'],
            'Proveedor': i['proveedor'],
            'Stock': i['stock'],
            'Precio Bs': i['precio_bs'],
            'Valor Total USD': i['valor_usd'],
            'Valor Total Bs': i['valor_bs'],
        } for i in data['items']])

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Libro Inventario')
            ws = writer.sheets['Libro Inventario']
            # Header rows before the table
            company = CompanySettingsService().get_company_context()
            ws.insert_rows(1, amount=5)
            ws['A1'] = company['company_name']
            ws['A2'] = f"RIF: {company['rif']}"
            ws['A3'] = f"Dirección Fiscal: {company['fiscal_address']}"
            ws['A4'] = f"Teléfono: {company['phone']}  |  Correo: {company['email']}"
            ws['A5'] = f"Fecha: {fecha.strftime('%d/%m/%Y')}  |  Tasa USD/Bs: {data['exchange_rate']}"
        output.seek(0)
        filename = f"libro_inventario_{fecha.strftime('%Y-%m-%d')}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.libro_inventario'))


# ---------------------------------------------------------------------------
# Libro Diario de Movimientos
# ---------------------------------------------------------------------------

MESES_ES = {
    1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril',
    5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto',
    9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre',
}


@reports_bp.route('/libro-diario')
@login_required
def libro_diario():
    """
    Libro diario de movimientos.
    Modos: 'hoy', 'dia', 'semana', 'semana_anterior', 'rango'.
    """
    try:
        modo = request.args.get('modo', 'hoy')
        if modo not in {'hoy', 'dia', 'semana', 'semana_anterior', 'rango'}:
            modo = 'hoy'
        today = date.today()

        if modo == 'hoy':
            start_date = end_date = today
        elif modo == 'dia':
            fecha_str = request.args.get('fecha', today.isoformat())
            try:
                start_date = end_date = datetime.strptime(fecha_str, '%Y-%m-%d').date()
            except ValueError:
                start_date = end_date = today
        elif modo == 'semana':
            # Lunes al domingo de la semana actual
            start_date = today - timedelta(days=today.weekday())
            end_date = start_date + timedelta(days=6)
            if end_date > today:
                end_date = today
        elif modo == 'semana_anterior':
            # Lunes al domingo de la semana pasada
            start_date = today - timedelta(days=today.weekday() + 7)
            end_date = start_date + timedelta(days=6)
        else:  # rango
            default_start = today - timedelta(days=6)
            start_str = request.args.get('start_date', default_start.isoformat())
            end_str = request.args.get('end_date', today.isoformat())
            try:
                start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
                end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
            except ValueError:
                start_date, end_date = today - timedelta(days=6), today

        service = ReportsService()
        data = service.get_libro_diario(start_date, end_date)

        return render_template(
            'reports/libro_diario.html',
            data=data,
            modo=modo,
            start_date=start_date,
            end_date=end_date,
            today=today,
        )
    except Exception as e:
        flash(f'Error al generar Libro Diario: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/libro-diario/export')
@login_required
def libro_diario_export():
    """Export libro diario to Excel."""
    try:
        today = date.today()
        start_str = request.args.get('start_date', today.isoformat())
        end_str = request.args.get('end_date', today.isoformat())
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d').date()
            end_date = datetime.strptime(end_str, '%Y-%m-%d').date()
        except ValueError:
            start_date = end_date = today

        service = ReportsService()
        data = service.get_libro_diario(start_date, end_date)

        # Sheet 1: journal entries in SENIAT visual layout
        lineas_rows = []
        for asiento in data['asientos']:
            for l in asiento['lineas']:
                es_autoconsumo = l['tipo'].upper() == 'SALIDA' and 'autoconsumo' in (l['descripcion_mov'] or '').lower()
                descripcion = (
                    f"{l['descripcion']} | {asiento['fecha'].strftime('%d/%m/%Y') if asiento['fecha'] else ''} "
                    f"{l['hora']} | {l['tipo']} | {l['descripcion_mov'] or 'Sin detalle'} | {l['usuario']}"
                )
                lineas_rows.append(_build_seniat_row(
                    codigo=l['codigo'],
                    descripcion=descripcion,
                    costo_inicial=l['precio_bs'],
                    costo_entrada=l['precio_bs'] if l['tipo'].upper() == 'ENTRADA' else 0,
                    cantidad_entrada=l['cantidad'] if l['tipo'].upper() == 'ENTRADA' else 0,
                    monto_entrada=l['valor_bs'] if l['tipo'].upper() == 'ENTRADA' else 0,
                    costo_salida=l['precio_bs'] if l['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                    cantidad_salida=l['cantidad'] if l['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                    monto_salida=l['valor_bs'] if l['tipo'].upper() == 'SALIDA' and not es_autoconsumo else 0,
                    costo_autoconsumo=l['precio_bs'] if es_autoconsumo else 0,
                    cantidad_autoconsumo=l['cantidad'] if es_autoconsumo else 0,
                    monto_autoconsumo=l['valor_bs'] if es_autoconsumo else 0,
                    costo_actual=l['precio_bs'],
                ))

        # Sheet 2: product summary in SENIAT visual layout
        resumen_rows = []
        for row in data['resumen_productos']:
            resumen_rows.append(_build_seniat_row(
                codigo=row['codigo'],
                descripcion=row['descripcion'],
                costo_inicial=_safe_unit_cost(row['valor_apertura_bs'], row['stock_apertura'], row.get('precio_bs_cierre', 0)),
                cantidad_inicial=row['stock_apertura'],
                monto_inicial=row['valor_apertura_bs'],
                costo_entrada=_safe_unit_cost(row.get('entradas_bs', 0), row['entradas'], row.get('precio_bs_cierre', 0)),
                cantidad_entrada=row['entradas'],
                monto_entrada=row.get('entradas_bs', 0),
                costo_salida=_safe_unit_cost(row.get('salidas_bs', 0), row['salidas'], row.get('precio_bs_cierre', 0)),
                cantidad_salida=row['salidas'],
                monto_salida=row.get('salidas_bs', 0),
                costo_autoconsumo=_safe_unit_cost(row.get('autoconsumos_bs', 0), row.get('autoconsumos', 0), row.get('precio_bs_cierre', 0)),
                cantidad_autoconsumo=row.get('autoconsumos', 0),
                monto_autoconsumo=row.get('autoconsumos_bs', 0),
                costo_actual=_safe_unit_cost(row['valor_cierre_bs'], row['stock_cierre'], row.get('precio_bs_cierre', 0)),
                cantidad_actual=row['stock_cierre'],
                monto_actual=row['valor_cierre_bs'],
            ))

        output = io.BytesIO()
        company = CompanySettingsService().get_company_context()
        workbook = _build_seniat_workbook([
            {
                'sheet_name': 'Asientos',
                'company': company,
                'start_date': start_date,
                'end_date': end_date,
                'rows': lineas_rows,
                'title': SENIAT_TEMPLATE_TITLE,
            },
            {
                'sheet_name': 'Resumen',
                'company': company,
                'start_date': start_date,
                'end_date': end_date,
                'rows': resumen_rows,
                'title': SENIAT_TEMPLATE_TITLE,
            },
        ])
        workbook.save(output)
        output.seek(0)
        filename = f"libro_diario_{start_date}_{end_date}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.libro_diario'))


# ---------------------------------------------------------------------------
# Resumen Mensual
# ---------------------------------------------------------------------------

@reports_bp.route('/resumen-mensual')
@login_required
def resumen_mensual():
    """Monthly inventory summary with opening/closing balances."""
    try:
        today = date.today()
        year = request.args.get('year', today.year, type=int)
        month = request.args.get('month', today.month, type=int)
        month = max(1, min(12, month))

        service = ReportsService()
        data = service.get_resumen_mensual(year, month)

        # Build year/month selectors
        years = list(range(today.year - 3, today.year + 1))

        return render_template(
            'reports/resumen_mensual.html',
            data=data,
            year=year,
            month=month,
            years=years,
            meses=MESES_ES,
            today=today,
        )
    except Exception as e:
        flash(f'Error al generar Resumen Mensual: {str(e)}', 'error')
        return redirect(url_for('reports.index'))


@reports_bp.route('/resumen-mensual/export')
@login_required
def resumen_mensual_export():
    """Export monthly summary to Excel."""
    try:
        today = date.today()
        year = request.args.get('year', today.year, type=int)
        month = request.args.get('month', today.month, type=int)
        month = max(1, min(12, month))

        service = ReportsService()
        data = service.get_resumen_mensual(year, month)

        filas = []
        for fila in data['filas']:
            entrada_monto = fila.get('entradas_bs', 0) or round(fila.get('precio_bs_apertura', 0) * fila['entradas'], 2)
            salida_monto = fila.get('salidas_bs', 0) or round(fila.get('precio_bs_apertura', 0) * fila['salidas'], 2)
            autoconsumo_monto = fila.get('autoconsumos_bs', 0) or round(fila.get('precio_bs_apertura', 0) * fila.get('autoconsumos', 0), 2)
            filas.append(_build_seniat_row(
                codigo=fila['codigo'],
                descripcion=fila['descripcion'],
                costo_inicial=fila.get('precio_bs_apertura', _safe_unit_cost(fila['valor_apertura_bs'], fila['stock_apertura'])),
                cantidad_inicial=fila['stock_apertura'],
                monto_inicial=fila['valor_apertura_bs'],
                costo_entrada=_safe_unit_cost(entrada_monto, fila['entradas'], fila.get('precio_bs_apertura', 0)),
                cantidad_entrada=fila['entradas'],
                monto_entrada=entrada_monto,
                costo_salida=_safe_unit_cost(salida_monto, fila['salidas'], fila.get('precio_bs_apertura', 0)),
                cantidad_salida=fila['salidas'],
                monto_salida=salida_monto,
                costo_autoconsumo=_safe_unit_cost(autoconsumo_monto, fila.get('autoconsumos', 0), fila.get('precio_bs_apertura', 0)),
                cantidad_autoconsumo=fila.get('autoconsumos', 0),
                monto_autoconsumo=autoconsumo_monto,
                costo_actual=fila.get('precio_bs_cierre', _safe_unit_cost(fila['valor_cierre_bs'], fila['stock_cierre'])),
                cantidad_actual=fila['stock_cierre'],
                monto_actual=fila['valor_cierre_bs'],
            ))

        output = io.BytesIO()
        company = CompanySettingsService().get_company_context()
        workbook = _build_seniat_workbook([{
            'sheet_name': 'Resumen Mensual',
            'company': company,
            'start_date': data['first_day'],
            'end_date': data['last_day'],
            'rows': filas,
            'title': SENIAT_TEMPLATE_TITLE,
        }])
        workbook.save(output)
        output.seek(0)
        filename = f"resumen_mensual_{year}_{month:02d}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        flash(f'Error al exportar: {str(e)}', 'error')
        return redirect(url_for('reports.resumen_mensual'))

